from __future__ import annotations

import io
import re

from app.schemas.lesson import Lesson, LessonSegment

MIME = {
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

INK = (42, 50, 56)
MUTED = (107, 115, 128)
ACCENT = (232, 163, 106)
CREAM = (255, 252, 248)
CARD_FILL = ((255, 244, 214), (229, 246, 234), (232, 243, 241))
CARD_INK = ((196, 138, 18), (47, 143, 91), (42, 125, 134))


def slugify(text: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", str(text or "").strip()).strip("-").lower()
    return slug[:60] or "training-deck"


def filename_for(lesson: Lesson, fmt: str) -> str:
    return f"{slugify(lesson.title)}.{fmt}"


def slide_points(segment: LessonSegment) -> tuple[str, str, list[str]]:
    slide = getattr(segment, "slide", None)
    headline = ""
    kicker = ""
    bullets: list[str] = []
    if slide is not None:
        headline = " ".join(str(getattr(slide, "headline", "") or "").split())
        kicker = " ".join(str(getattr(slide, "kicker", "") or "").split())
        for item in getattr(slide, "bullets", None) or []:
            text = " ".join(str(item or "").split())
            if len(text.split()) < 3:
                continue
            bullets.append(text)
            if len(bullets) == 3:
                break
    if not headline:
        headline = " ".join(str(segment.title or "").split()) or "Key idea"
    if re.match(r"^line\s+\d+$", headline, re.I):
        headline = "Key idea"
    if not kicker:
        kicker = "Live training demo"
    if not bullets:
        speech = " ".join(str(segment.speech or "").split())
        if speech:
            bullets = [speech[:110] + ("…" if len(speech) > 110 else "")]
    return headline, kicker, bullets[:3]


def build_export(lesson: Lesson, fmt: str, brand: str = "") -> tuple[bytes, str, str]:
    kind = str(fmt or "").lower().strip()
    if kind not in MIME:
        raise ValueError("format must be pptx, pdf, or docx")
    if not lesson.segments:
        raise ValueError("lesson has no slides")
    builders = {"pptx": _pptx, "pdf": _pdf, "docx": _docx}
    data = builders[kind](lesson, brand.strip())
    return data, MIME[kind], filename_for(lesson, kind)


def _pptx(lesson: Lesson, brand: str) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    blank = deck.slide_layouts[6]
    ink = RGBColor(*INK)
    muted = RGBColor(*MUTED)
    cream = RGBColor(*CREAM)
    accent = RGBColor(*ACCENT)

    def paint(slide) -> None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = cream
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), deck.slide_width, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()
        if brand:
            box = slide.shapes.add_textbox(Inches(0.7), Inches(7.05), Inches(8), Inches(0.28))
            tf = box.text_frame
            tf.text = brand
            p = tf.paragraphs[0]
            p.font.size = Pt(11)
            p.font.color.rgb = muted
            p.font.bold = True

    def heading(slide, text: str, top: float, size: int = 36, width: float = 12.0) -> None:
        box = slide.shapes.add_textbox(Inches(0.7), Inches(top), Inches(width), Inches(1.15))
        tf = box.text_frame
        tf.word_wrap = True
        tf.text = text
        p = tf.paragraphs[0]
        p.font.size = Pt(size)
        p.font.color.rgb = ink
        p.font.bold = True
        p.alignment = PP_ALIGN.LEFT

    def eyebrow(slide, text: str, top: float) -> None:
        box = slide.shapes.add_textbox(Inches(0.7), Inches(top), Inches(12), Inches(0.32))
        tf = box.text_frame
        tf.text = text.upper()
        p = tf.paragraphs[0]
        p.font.size = Pt(12)
        p.font.color.rgb = accent
        p.font.bold = True

    def cards(slide, bullets: list[str]) -> None:
        count = max(len(bullets), 1)
        gap = 0.28
        left = 0.7
        usable = 13.333 - (left * 2) - gap * (count - 1)
        width = usable / count
        top = 3.05
        height = 3.35
        for index, bullet in enumerate(bullets):
            x = left + index * (width + gap)
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(x),
                Inches(top),
                Inches(width),
                Inches(height),
            )
            try:
                shape.adjustments[0] = 0.12
            except Exception:
                pass
            fill = CARD_FILL[index % 3]
            tint = CARD_INK[index % 3]
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(*fill)
            shape.line.color.rgb = RGBColor(*fill)
            tf = shape.text_frame
            tf.word_wrap = True
            tf.margin_left = Emu(140000)
            tf.margin_right = Emu(140000)
            tf.margin_top = Emu(160000)
            tf.margin_bottom = Emu(120000)
            tf.text = f"{index + 1:02d}"
            num = tf.paragraphs[0]
            num.font.size = Pt(14)
            num.font.bold = True
            num.font.color.rgb = RGBColor(*tint)
            body = tf.add_paragraph()
            body.text = bullet
            body.font.size = Pt(18)
            body.font.color.rgb = ink
            body.space_before = Pt(12)
            body.font.bold = False

    def spoken_notes(slide, text: str) -> None:
        if not text:
            return
        notes = slide.notes_slide
        notes.notes_text_frame.text = text

    title = deck.slides.add_slide(blank)
    paint(title)
    eyebrow(title, brand or "Training demo", 2.05)
    heading(title, lesson.title or "Training", 2.45, 48)
    sub = title.shapes.add_textbox(Inches(0.7), Inches(3.8), Inches(11), Inches(0.8))
    tf = sub.text_frame
    tf.word_wrap = True
    tf.text = "The same slides shown in the live classroom."
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.color.rgb = muted

    total = len(lesson.segments)
    for index, segment in enumerate(lesson.segments, start=1):
        headline, kicker, bullets = slide_points(segment)
        slide = deck.slides.add_slide(blank)
        paint(slide)
        eyebrow(slide, f"{kicker}  ·  {index} / {total}", 0.42)
        heading(slide, headline, 0.78, 32)
        cards(slide, bullets)
        try:
            spoken_notes(slide, segment.speech)
        except Exception:
            pass
        counter = slide.shapes.add_textbox(Inches(11.2), Inches(7.05), Inches(1.4), Inches(0.28))
        ctf = counter.text_frame
        ctf.text = f"{index:02d} / {total:02d}"
        ctf.paragraphs[0].font.size = Pt(11)
        ctf.paragraphs[0].font.color.rgb = muted
        ctf.paragraphs[0].alignment = PP_ALIGN.RIGHT

    thanks = deck.slides.add_slide(blank)
    paint(thanks)
    eyebrow(thanks, brand or "Training complete", 2.2)
    heading(thanks, "Thank you", 2.6, 52)
    body = thanks.shapes.add_textbox(Inches(0.7), Inches(3.95), Inches(11), Inches(0.6))
    btf = body.text_frame
    btf.text = "Download this deck as PowerPoint, PDF, or Word."
    btf.paragraphs[0].font.size = Pt(18)
    btf.paragraphs[0].font.color.rgb = muted

    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def _pdf(lesson: Lesson, brand: str) -> bytes:
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase.pdfmetrics import stringWidth

    buffer = io.BytesIO()
    width, height = landscape(A4)
    pdf = canvas.Canvas(buffer, pagesize=landscape(A4))

    def rgb(color: tuple[int, int, int]) -> tuple[float, float, float]:
        return color[0] / 255, color[1] / 255, color[2] / 255

    def wrapped(text: str, font: str, size: int, max_width: float) -> list[str]:
        words = str(text or "").split()
        lines: list[str] = []
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if stringWidth(candidate, font, size) <= max_width:
                current = candidate
                continue
            if current:
                lines.append(current)
            current = word
        if current:
            lines.append(current)
        return lines or [""]

    def chrome(kicker: str, footer: str = "") -> None:
        pdf.setFillColorRGB(*rgb(CREAM))
        pdf.rect(0, 0, width, height, fill=1, stroke=0)
        pdf.setFillColorRGB(*rgb(ACCENT))
        pdf.rect(0, height - 8, width, 8, fill=1, stroke=0)
        pdf.setFillColorRGB(*rgb(ACCENT))
        pdf.setFont("Helvetica-Bold", 11)
        pdf.drawString(22 * mm, height - 18 * mm, (kicker or "LIVE TRAINING DEMO").upper()[:72])
        if brand:
            pdf.setFillColorRGB(*rgb(MUTED))
            pdf.setFont("Helvetica-Bold", 10)
            pdf.drawString(22 * mm, 12 * mm, brand)
        if footer:
            pdf.setFillColorRGB(*rgb(MUTED))
            pdf.setFont("Helvetica", 10)
            pdf.drawRightString(width - 22 * mm, 12 * mm, footer)

    def title_block(title: str, top: float, size: int = 28) -> float:
        pdf.setFillColorRGB(*rgb(INK))
        pdf.setFont("Times-Bold", size)
        y = top
        for line in wrapped(title, "Times-Bold", size, width - 48 * mm)[:3]:
            pdf.drawString(22 * mm, y, line)
            y -= size * 1.15
        return y

    def draw_cards(bullets: list[str], top: float) -> None:
        count = max(len(bullets), 1)
        gap = 8 * mm
        left = 22 * mm
        usable = width - 44 * mm - gap * (count - 1)
        card_w = usable / count
        card_h = 72 * mm
        radius = 10
        for index, bullet in enumerate(bullets):
            x = left + index * (card_w + gap)
            y = top - card_h
            fill = CARD_FILL[index % 3]
            tint = CARD_INK[index % 3]
            pdf.setFillColorRGB(*rgb(fill))
            pdf.roundRect(x, y, card_w, card_h, radius, fill=1, stroke=0)
            pdf.setFillColorRGB(*rgb(tint))
            pdf.setFont("Helvetica-Bold", 12)
            pdf.drawString(x + 8 * mm, y + card_h - 12 * mm, f"{index + 1:02d}")
            pdf.setFillColorRGB(*rgb(INK))
            pdf.setFont("Helvetica", 13)
            text = pdf.beginText(x + 8 * mm, y + card_h - 24 * mm)
            text.setLeading(18)
            for line in wrapped(bullet, "Helvetica", 13, card_w - 16 * mm)[:5]:
                text.textLine(line)
            pdf.drawText(text)

    chrome(brand or "Training demo")
    title_block(lesson.title or "Training", height - 42 * mm, 36)
    pdf.setFillColorRGB(*rgb(MUTED))
    pdf.setFont("Helvetica", 16)
    pdf.drawString(22 * mm, height - 78 * mm, "The same slides shown in the live classroom.")
    pdf.showPage()

    total = len(lesson.segments)
    for index, segment in enumerate(lesson.segments, start=1):
        headline, kicker, bullets = slide_points(segment)
        chrome(kicker, f"{index:02d} / {total:02d}")
        y = title_block(headline, height - 34 * mm, 26)
        draw_cards(bullets, y - 8 * mm)
        pdf.showPage()

    chrome(brand or "Training complete")
    title_block("Thank you", height - 50 * mm, 42)
    pdf.setFillColorRGB(*rgb(MUTED))
    pdf.setFont("Helvetica", 16)
    pdf.drawString(22 * mm, height - 82 * mm, "Download this deck as PowerPoint, PDF, or Word.")
    pdf.showPage()

    pdf.save()
    return buffer.getvalue()


def _docx(lesson: Lesson, brand: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    if brand:
        tag = doc.add_paragraph(brand)
        tag.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if tag.runs:
            tag.runs[0].font.size = Pt(11)
            tag.runs[0].font.color.rgb = RGBColor(*MUTED)
            tag.runs[0].bold = True

    title = doc.add_heading(lesson.title or "Training", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(*INK)

    intro = doc.add_paragraph("The same slides shown in the live classroom.")
    if intro.runs:
        intro.runs[0].font.size = Pt(12)
        intro.runs[0].font.color.rgb = RGBColor(*MUTED)

    for index, segment in enumerate(lesson.segments, start=1):
        headline, kicker, bullets = slide_points(segment)
        heading = doc.add_heading(f"{index}. {headline}", level=1)
        for run in heading.runs:
            run.font.color.rgb = RGBColor(*INK)
        if kicker:
            eye = doc.add_paragraph(kicker)
            if eye.runs:
                eye.runs[0].font.size = Pt(11)
                eye.runs[0].font.color.rgb = RGBColor(*ACCENT)
                eye.runs[0].bold = True
        for bullet in bullets:
            doc.add_paragraph(bullet, style="List Bullet")
        if segment.speech:
            spoken = doc.add_paragraph(segment.speech)
            if spoken.runs:
                spoken.runs[0].italic = True
                spoken.runs[0].font.color.rgb = RGBColor(*MUTED)

    doc.add_heading("Thank you", level=1)
    doc.add_paragraph("This document matches the presented training deck.")

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
